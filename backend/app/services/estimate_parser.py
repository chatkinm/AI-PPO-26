import openpyxl
from typing import List, Dict, Any
from pathlib import Path

class EstimateParser:
    """Парсер Excel-смет Xpage"""

    def parse_file(self, file_path: Path) -> Dict[str, Any]:
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        sheet_names = wb.sheetnames
        sheet = wb.active
        if "Актуальная смета Приложение" in sheet_names:
            sheet = wb["Актуальная смета Приложение"]

        rows_data: List[Dict[str, Any]] = []
        current_section = "Общее"

        for r in range(1, sheet.max_row + 1):
            col1 = sheet.cell(r, 1).value
            col2 = sheet.cell(r, 2).value
            
            val = str(col1 or col2 or "").strip()
            if not val:
                continue

            if len(val) < 60 and not "—" in val and not val.startswith("-") and not val.replace('.', '').isdigit():
                current_section = val
                continue

            lines = val.split("\n")
            sub_items = [line.strip().lstrip("— -").strip() for line in lines if line.strip().lstrip("— -")]
            title = sub_items[0] if sub_items else val
            elements = sub_items[1:] if len(sub_items) > 1 else []

            hours = 0.0
            for c in range(2, min(sheet.max_column + 1, 10)):
                cell_val = sheet.cell(r, c).value
                if isinstance(cell_val, (int, float)) and cell_val > 0:
                    hours += float(cell_val)

            rows_data.append({
                "row_index": r,
                "section": current_section,
                "raw_text": val,
                "title": title,
                "elements": elements,
                "hours": hours
            })

        return {
            "sheet_name": sheet.title,
            "total_rows": len(rows_data),
            "rows": rows_data
        }

estimate_parser = EstimateParser()
